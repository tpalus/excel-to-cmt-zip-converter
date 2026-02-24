"""
Excel file loading and processing.
"""

from pathlib import Path
from typing import Dict, List, Optional

import pandas as pd
from openpyxl import load_workbook


class ExcelLoader:
    """Loads data from Excel files into DataFrames."""

    def __init__(self, excel_path: Path):
        """
        Initialize Excel loader.
        
        Args:
            excel_path: Path to Excel file
        """
        if not excel_path.exists():
            raise FileNotFoundError(f"Excel file not found: {excel_path}")

        self.excel_path = excel_path

    def load_all_tables(self) -> Dict[str, pd.DataFrame]:
        """
        Load all tables from Excel workbook.
        
        Returns:
            Dictionary mapping table names to DataFrames
        """
        wb = load_workbook(str(self.excel_path), data_only=True, read_only=False)
        tables = {}

        try:
            for worksheet in wb.worksheets:
                for table_name, table in worksheet.tables.items():
                    ref = table.ref if hasattr(table, 'ref') else table
                    data = worksheet[ref]
                    rows = [[cell.value for cell in row] for row in data]

                    if not rows:
                        continue

                    header = rows[0]
                    body = rows[1:]

                    tables[table_name] = pd.DataFrame(body, columns=header)

        finally:
            wb.close()

        return tables

    def filter_tables(
        self,
        tables: Dict[str, pd.DataFrame],
        columns_to_keep: Dict[str, List[str]],
        from_safe_str_func
    ) -> Dict[str, pd.DataFrame]:
        """
        Filter tables based on columns_to_keep specification.
        
        Args:
            tables: Dictionary of DataFrames
            columns_to_keep: Dict mapping entity names to list of columns
                           Empty dict {} = include all entities with all columns
                           Empty list [] = include entity with all columns
                           Columns starting with '-' = exclude those
                           Otherwise = include specified columns only
            from_safe_str_func: Function to convert values to strings
            
        Returns:
            Filtered dictionary of DataFrames
        """
        filtered = {}
        include_all = not columns_to_keep  # True if dict is empty

        for entity_name, df in tables.items():
            # Skip only if columns_to_keep is specified AND entity not in it
            if not include_all and entity_name not in columns_to_keep:
                print(f"Skipping entity '{entity_name}' not in columns_to_keep")
                continue

            pk_col = f"{entity_name}id"
            
            # Get specification for this entity
            spec = columns_to_keep.get(entity_name, []) if columns_to_keep else []

            # Determine which columns to keep
            if not spec:
                # No restriction - keep all
                cols = list(df.columns)
            elif all(col.startswith('-') for col in spec):
                # All specifications are negative - exclude these
                exclude = [col[1:] for col in spec]
                cols = [c for c in df.columns if c not in exclude]
            else:
                # Positive selection - keep only these
                include = [col for col in spec if not col.startswith('-')]
                cols = [c for c in include if c in df.columns]

            # Add primary key if not already included
            if pk_col in df.columns and pk_col not in cols:
                cols.append(pk_col)

            if not cols:
                print(f"No columns found for entity '{entity_name}' with spec {spec}")
                continue

            filtered_df = df[cols].copy()

            # Convert all values to strings
            for column in filtered_df.columns:
                filtered_df[column] = filtered_df[column].apply(from_safe_str_func)

            filtered[entity_name] = filtered_df

        return filtered
